# control_server.py  —— 实时会话 + 事件流 + 导出
import os, asyncio, json, uuid, re, logging
from typing import Optional, Dict, Any, List

from fastapi import FastAPI, Body
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import PlainTextResponse, StreamingResponse, JSONResponse, FileResponse

from client import MCPClient  # 复用你的 MCP 客户端

LOG_PATH = os.getenv("LOG_PATH", "test.log")
EXPORT_DIR = os.getenv("EXPORT_DIR", "outputs")
os.makedirs(EXPORT_DIR, exist_ok=True)

app = FastAPI(title="MCP Controller", version="2.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"], allow_credentials=False, allow_methods=["*"], allow_headers=["*"],
)

logger = logging.getLogger("control_server")
if not logger.handlers:
    logging.basicConfig(level=logging.INFO)

_client: Optional[MCPClient] = None
_client_lock = asyncio.Lock()
_active_choice = {
    "requested_provider": None,
    "active_provider": None,
    "requested_model": None,
    "resolved_model": None,
}

class LiveSession:
    def __init__(self, sid:str, provider: Optional[str] = None, model: Optional[str] = None):
        self.id = sid
        self.queue: asyncio.Queue[str] = asyncio.Queue()
        self.task: Optional[asyncio.Task] = None
        self.active: bool = False
        self.final_md: str = ""
        self.graph_nodes: List[Dict[str,Any]] = []
        self.graph_edges: List[Dict[str,Any]] = []
        self.provider = provider
        self.model = model

_sessions: Dict[str, LiveSession] = {}

# --------- 工具函数 ---------
async def _start_client_locked(provider_override: Optional[str], model_override: Optional[str]):
    global _client, _active_choice
    if _client:
        await _client.close()
    client = MCPClient(provider_override=provider_override, model_override=model_override)
    await client.connect_to_server("./search_mcp.py")
    _client = client
    _active_choice["requested_provider"] = provider_override or getattr(client.llm, "provider_mode", None)
    _active_choice["requested_model"] = model_override or getattr(client.llm, "requested_model", None)
    _active_choice["active_provider"] = getattr(client.llm, "active_provider", None) or provider_override
    _active_choice["resolved_model"] = getattr(client.llm, "model_name", None)
    logger.info(
        "[Controller] MCP client active provider=%s requested=%s resolved=%s",
        _active_choice["active_provider"],
        model_override,
        _active_choice["resolved_model"],
    )


async def ensure_client_started(provider: Optional[str] = None, model: Optional[str] = None):
    global _client
    async with _client_lock:
        if _client is None:
            await _start_client_locked(
                provider if provider is not None else _active_choice.get("requested_provider"),
                model if model is not None else _active_choice.get("requested_model"),
            )
        else:
            if provider and provider != _active_choice.get("requested_provider"):
                logger.info(
                    "[Controller] client already running with provider=%s, ignore requested provider=%s",
                    _active_choice.get("requested_provider"),
                    provider,
                )
            if model and model != _active_choice.get("requested_model"):
                logger.info(
                    "[Controller] client already running with model=%s, ignore requested model=%s",
                    _active_choice.get("requested_model"),
                    model,
                )

def sse_format(data: dict) -> str:
    return f"data: {json.dumps(data, ensure_ascii=False)}\n\n"

# --------- 控制器原有 ---------
@app.post("/api/start")
async def start_controller(payload: Optional[dict] = Body(default=None)):
    data = payload or {}
    provider = data.get("provider")
    model = data.get("model")
    async with _client_lock:
        await _start_client_locked(provider, model)
    logger.info("[Controller] start requested provider=%s model=%s", provider, model)
    return {
        "status": "started",
        "provider": _active_choice.get("active_provider"),
        "model": _active_choice.get("resolved_model"),
    }

@app.post("/api/stop")
async def stop_controller():
    global _client
    async with _client_lock:
        if _client:
            await _client.close()
        _client = None
        _active_choice.update({
            "active_provider": None,
            "requested_provider": None,
            "requested_model": None,
            "resolved_model": None,
        })
    return {"status":"stopped"}

@app.get("/api/status")
async def status():
    running = _client is not None
    provider = _active_choice.get("active_provider")
    model = _active_choice.get("resolved_model")
    requested_model = _active_choice.get("requested_model")
    ollama_models = _client.llm.ollama.models if (_client and _client.llm and _client.llm.ollama) else []
    return {
        "running": bool(running),
        "provider": provider,
        "model": model,
        "requested_model": requested_model,
        "ollama_models": ollama_models,
    }

@app.get("/api/logs/tail", response_class=PlainTextResponse)
async def logs_tail(lines: int = 400):
    try:
        with open(LOG_PATH, "r", encoding="utf-8", errors="ignore") as f:
            data = f.readlines()
        return "".join(data[-abs(int(lines)):])
    except FileNotFoundError:
        return ""
    except Exception as e:
        return PlainTextResponse(f"[error reading log] {e}", status_code=500)

# --------- 新：会话 / 事件流 / 交互 ----------
@app.post("/api/session")
async def create_session(payload: Optional[dict] = Body(default=None)):
    data = payload or {}
    sid = uuid.uuid4().hex
    sess = LiveSession(sid, provider=data.get("provider"), model=data.get("model"))
    _sessions[sid] = sess
    logger.info("[Session:%s] created with model=%s provider=%s", sid, sess.model, sess.provider)
    return {"session_id": sid, "model": sess.model, "provider": sess.provider}

@app.post("/api/session/{sid}/clarify")
async def clarify(sid: str, payload: dict = Body(...)):
    sess = _sessions.get(sid)
    model = payload.get("model") if isinstance(payload, dict) else None
    provider = payload.get("provider") if isinstance(payload, dict) else None
    if sess:
        if model:
            sess.model = model
        if provider:
            sess.provider = provider
        provider = provider or sess.provider
        model = model or sess.model
    await ensure_client_started(provider=provider, model=model)
    question = (payload.get("question") or "").strip()
    if not question:
        return JSONResponse({"error":"empty_question"}, status_code=400)
    prompt = f"请针对下面的用户问题提出 3-5 条澄清问题，编号列出，尽量覆盖范围、目标、限制、偏好等关键信息。\n\n用户问题：{question}"
    resp = await _client.llm.create_chat_completion([
        {"role":"system","content":"你是严谨的研究助理。只输出问题列表，不要多余说明。"},
        {"role":"user","content": prompt}
    ])
    text = resp.choices[0].message.content
    questions = re.split(r'^\s*\d+[\.\)]\s*', text, flags=re.M)[1:] or [text]
    return {"clarifying": [q.strip() for q in questions if q.strip()]}

@app.post("/api/session/{sid}/start")
async def start_research(sid:str, payload: dict = Body(...)):
    sess = _sessions.get(sid)
    provider = payload.get("provider") if isinstance(payload, dict) else None
    model = payload.get("model") if isinstance(payload, dict) else None
    if sess:
        if provider:
            sess.provider = provider
        if model:
            sess.model = model
        provider = provider or sess.provider
        model = model or sess.model
    await ensure_client_started(provider=provider, model=model)
    if sess:
        logger.info(
            "[Session:%s] starting research with requested_model=%s active_model=%s",
            sid,
            sess.model,
            _active_choice.get("resolved_model"),
        )
    else:
        logger.info(
            "[Session:%s] starting research with active_model=%s",
            sid,
            _active_choice.get("resolved_model"),
        )
    sess = _sessions.get(sid)
    if not sess: return JSONResponse({"error":"session_not_found"}, status_code=404)

    query = (payload.get("query") or "").strip()
    if not query: return JSONResponse({"error":"empty_query"}, status_code=400)

    async def sink(ev: dict):
        await sess.queue.put(sse_format(ev))
        t = ev.get("type")
        if t == "tool_call":
            nid = f"call_{len(sess.graph_nodes)}"
            sess.graph_nodes.append({"id": nid, "label": f"🔧 {ev.get('name')}"})
            if sess.graph_nodes:
                prev = sess.graph_nodes[-2]["id"] if len(sess.graph_nodes) >= 2 else "root"
                sess.graph_edges.append({"source": prev, "target": nid})
        elif t == "tool_result":
            nid = f"res_{len(sess.graph_nodes)}"
            sess.graph_nodes.append({"id": nid, "label": f"📄 {ev.get('name')} 结果"})
            prev = sess.graph_nodes[-2]["id"] if len(sess.graph_nodes) >= 2 else "root"
            sess.graph_edges.append({"source": prev, "target": nid})
        elif t == "final_report" and ev.get("format") == "markdown":
            sess.final_md = ev.get("content","")

    async def runner():
        sess.active = True
        await sess.queue.put(sse_format({"type":"phase","phase":"research","progress":0.15}))
        try:
            await _client.process_query_stream(query, event_cb=sink)
        except Exception as e:
            await sess.queue.put(sse_format({"type":"error","message":str(e)}))
        finally:
            sess.active = False
            await sess.queue.put(sse_format({"type":"phase","phase":"done","progress":1.0}))

    if sess.task and not sess.task.done():
        return {"status":"already_running"}
    sess.task = asyncio.create_task(runner())
    return {"status":"started"}

@app.get("/api/session/{sid}/events/stream")
async def events_stream(sid: str):
    sess = _sessions.get(sid)
    if not sess: return JSONResponse({"error":"session_not_found"}, status_code=404)

    async def gen():
        yield sse_format({"type":"graph_init","root":{"id":"root","label":"🧭 研究会话"}})
        while True:
            msg = await sess.queue.get()
            yield msg
    return StreamingResponse(gen(), media_type="text/event-stream")

@app.get("/api/session/{sid}/report")
async def get_report(sid:str):
    sess = _sessions.get(sid)
    if not sess: return JSONResponse({"error":"session_not_found"}, status_code=404)
    return {"markdown": sess.final_md or ""}

# --------- 导出（DOCX / PDF / Markdown） ----------
def export_docx(md: str, filepath: str):
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    lines = md.splitlines()
    in_code = False
    buf: List[str] = []

    def flush_buf():
        nonlocal buf
        if buf:
            doc.add_paragraph("\n".join(buf))
            buf.clear()

    for ln in lines:
        if ln.strip().startswith("```"):
            if in_code:
                doc.add_paragraph("\n".join(buf)).style = 'Intense Quote'
                buf.clear(); in_code=False
            else:
                flush_buf(); in_code=True
            continue
        if in_code:
            buf.append(ln); continue
        if ln.startswith("# "): flush_buf(); doc.add_heading(ln[2:].strip(), level=1); continue
        if ln.startswith("## "): flush_buf(); doc.add_heading(ln[3:].strip(), level=2); continue
        if ln.startswith("### "): flush_buf(); doc.add_heading(ln[4:].strip(), level=3); continue
        if re.match(r"^\s*[-*]\s+", ln):
            doc.add_paragraph(re.sub(r"^\s*[-*]\s+","",ln), style='List Bullet'); continue
        if re.match(r"^\s*\d+\.\s+", ln):
            doc.add_paragraph(re.sub(r"^\s*\d+\.\s+","",ln), style='List Number'); continue
        doc.add_paragraph(ln)
    flush_buf()
    doc.save(filepath)

def export_pdf_simple(md: str, filepath: str):
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    font_path = os.getenv("PDF_FONT_TTF")
    if font_path and os.path.exists(font_path):
        pdfmetrics.registerFont(TTFont("DocFont", font_path))
        base_font = "DocFont"
    else:
        base_font = "Helvetica"

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="Body", fontName=base_font, fontSize=10.5, leading=14))
    styles.add(ParagraphStyle(name="H1", parent=styles['Heading1'], fontName=base_font))
    styles.add(ParagraphStyle(name="H2", parent=styles['Heading2'], fontName=base_font))
    styles.add(ParagraphStyle(name="H3", parent=styles['Heading3'], fontName=base_font))

    doc = SimpleDocTemplate(filepath, pagesize=A4)
    story: List[Any] = []
    lines = md.splitlines()
    in_code, codebuf = False, []

    def flush_code():
        nonlocal codebuf
        if codebuf:
            story.append(Paragraph("<br/>".join([x.replace("<","&lt;").replace(">","&gt;") for x in codebuf]), styles["Body"]))
            story.append(Spacer(1,6))
            codebuf=[]
    for ln in lines:
        if ln.strip().startswith("```"):
            if in_code: flush_code(); in_code=False
            else: in_code=True
            continue
        if in_code:
            codebuf.append(ln); continue
        if ln.startswith("# "): story.append(Paragraph(ln[2:].strip(), styles["H1"])); story.append(Spacer(1,8)); continue
        if ln.startswith("## "): story.append(Paragraph(ln[3:].strip(), styles["H2"])); story.append(Spacer(1,6)); continue
        if ln.startswith("### "): story.append(Paragraph(ln[4:].strip(), styles["H3"])); story.append(Spacer(1,4)); continue
        if re.match(r"^\s*[-*]\s+", ln):
            story.append(Paragraph("• "+re.sub(r"^\s*[-*]\s+","",ln), styles["Body"])); continue
        story.append(Paragraph(ln.replace("  ", "&nbsp;&nbsp;"), styles["Body"]))
    flush_code()
    doc.build(story)

@app.post("/api/session/{sid}/export")
async def export_report(sid: str, payload: dict = Body(...)):
    sess = _sessions.get(sid)
    if not sess: return JSONResponse({"error":"session_not_found"}, status_code=404)
    md = (sess.final_md or "").strip()
    if not md: return JSONResponse({"error":"no_report"}, status_code=400)

    fmt = (payload.get("format") or "md").lower()
    base = os.path.join(EXPORT_DIR, f"deepresearch_{sid}")
    if fmt == "md":
        path = base + ".md"
        with open(path, "w", encoding="utf-8") as f: f.write(md)
        return FileResponse(path, filename=os.path.basename(path), media_type="text/markdown")
    elif fmt == "docx":
        path = base + ".docx"
        export_docx(md, path)
        return FileResponse(path, filename=os.path.basename(path), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    elif fmt == "pdf":
        path = base + ".pdf"
        export_pdf_simple(md, path)
        return FileResponse(path, filename=os.path.basename(path), media_type="application/pdf")
    return JSONResponse({"error":"unsupported_format"}, status_code=400)

# 允许直接 python control_server.py 启动
if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="127.0.0.1", port=8787)
